import streamlit as st
import anthropic
import requests
import json
import time
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os

# ── Text sanitizer ────────────────────────────────────────────────
def clean(text):
    """Remove non-ASCII and problematic Unicode characters from text."""
    if not text:
        return ""
    if not isinstance(text, str):
        text = str(text)
    # Replace common problematic unicode chars
    replacements = {
        ' ': ' ',  # Line separator
        ' ': ' ',  # Paragraph separator
        '​': '',   # Zero width space
        '‌': '',   # Zero width non-joiner
        '‍': '',   # Zero width joiner
        '﻿': '',   # BOM
        ' ': ' ',  # Non-breaking space
        '–': '-',  # En dash
        '—': '-',  # Em dash
        '‘': "'",  # Left single quote
        '’': "'",  # Right single quote
        '“': '"',  # Left double quote
        '”': '"',  # Right double quote
        '…': '...', # Ellipsis
    }
    for char, replacement in replacements.items():
        text = text.replace(char, replacement)
    # Final fallback: encode to ASCII ignoring errors
    return text.encode('ascii', 'ignore').decode('ascii')

# ── Page config ────────────────────────────────────────────────────
st.set_page_config(
    page_title="Daily Job Pipeline",
    page_icon="🚀",
    layout="wide"
)

st.markdown("""
<style>
.stage-box {
    padding: 12px 16px;
    border-radius: 8px;
    border: 1px solid #e0e0e0;
    margin-bottom: 8px;
    background: #fafafa;
}
.stage-done { border-color: #4caf50; background: #f0fff0; }
.stage-active { border-color: #2196f3; background: #f0f8ff; }
.stage-pending { border-color: #e0e0e0; background: #fafafa; }
.job-card {
    padding: 14px;
    border-radius: 8px;
    border: 1px solid #e0e0e0;
    margin-bottom: 10px;
    cursor: pointer;
}
.score-high { color: #4caf50; font-weight: 700; }
.score-med { color: #2196f3; font-weight: 700; }
.email-found { color: #4caf50; font-family: monospace; font-size: 13px; }
.email-miss { color: #999; font-style: italic; font-size: 12px; }
</style>
""", unsafe_allow_html=True)

# ── Constants ──────────────────────────────────────────────────────
APIFY_BASE = "https://api.apify.com/v2"
HUNTER_BASE = "https://api.hunter.io/v2"

SEARCH_QUERIES = [
    "Senior Product Manager Authentication Identity Security AI enterprise remote 2026",
    "Senior Product Manager API Platform LLM Generative AI B2B SaaS 2026",
    "Director Product Manager Enterprise AI Security Compliance Platform 2026",
    "Principal Product Manager Platform Identity Security Cloud 2026",
    "Senior PM Fintech AI Products Authentication enterprise 2026"
]

# ── API helpers ────────────────────────────────────────────────────

def run_apify(actor: str, input_data: dict, apify_key: str, timeout: int = 300) -> list:
    """Run an Apify actor and return results."""
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {apify_key}"
    }
    # Start run
    resp = requests.post(
        f"{APIFY_BASE}/acts/{actor}/runs",
        headers=headers,
        json=input_data,
        timeout=30
    )
    resp.raise_for_status()
    run_id = resp.json()["data"]["id"]

    # Poll until done
    deadline = time.time() + timeout
    while time.time() < deadline:
        time.sleep(5)
        status_resp = requests.get(
            f"{APIFY_BASE}/actor-runs/{run_id}",
            headers=headers,
            timeout=15
        )
        status = status_resp.json()["data"]["status"]
        if status == "SUCCEEDED":
            dataset_id = status_resp.json()["data"]["defaultDatasetId"]
            items_resp = requests.get(
                f"{APIFY_BASE}/datasets/{dataset_id}/items?limit=100",
                headers=headers,
                timeout=30
            )
            return items_resp.json()
        if status in ("FAILED", "ABORTED"):
            raise RuntimeError(f"Apify run {status}")
    raise TimeoutError("Apify run timed out")


def call_claude(system: str, user: str, api_key: str) -> str:
    """Call Claude API and return text response."""
    client = anthropic.Anthropic(api_key=api_key)
    msg = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=1000,
        system=system,
        messages=[{"role": "user", "content": user}]
    )
    return msg.content[0].text


def parse_json_safe(text: str):
    """Parse JSON safely, stripping code fences."""
    try:
        clean = text.replace("```json", "").replace("```", "").strip()
        return json.loads(clean)
    except Exception:
        return None


def hunter_find_email(first: str, last: str, company: str, hunter_key: str) -> dict | None:
    """Find email using Hunter.io."""
    if not hunter_key or not first or not last:
        return None
    try:
        resp = requests.get(
            f"{HUNTER_BASE}/email-finder",
            params={
                "company": company,
                "first_name": first,
                "last_name": last,
                "api_key": hunter_key
            },
            timeout=10
        )
        data = resp.json()
        if data.get("data", {}).get("email"):
            return {
                "email": data["data"]["email"],
                "score": data["data"].get("score", 0),
                "verified": data["data"].get("score", 0) > 70
            }
    except Exception:
        pass
    return None


def parse_name(full_name: str) -> tuple[str, str]:
    """Split full name into first and last."""
    parts = (full_name or "").strip().split()
    return parts[0] if parts else "", " ".join(parts[1:]) if len(parts) > 1 else ""


# ── Pipeline stages ────────────────────────────────────────────────

def stage_discover(apify_key: str, status_placeholder) -> list:
    status_placeholder.info("🔍 Searching Greenhouse, Ashby, Lever, LinkedIn, Google Jobs...")
    from concurrent.futures import ThreadPoolExecutor, as_completed

    def search_query(q):
        return run_apify(
            "apify/rag-web-browser",
            {"query": q, "maxResults": 5},
            apify_key,
            timeout=180
        )

    all_results = []
    with ThreadPoolExecutor(max_workers=3) as ex:
        futures = {ex.submit(search_query, q): q for q in SEARCH_QUERIES[:4]}
        for i, future in enumerate(as_completed(futures)):
            status_placeholder.info(f"🔍 Scraping job boards... ({i+1}/{len(futures)} queries complete)")
            try:
                all_results.extend(future.result() or [])
            except Exception as e:
                pass

    content = "\n---\n".join([
        f"SOURCE: {r.get('url','')}\nTITLE: {r.get('title','')}\nCONTENT: {str(r.get('text',''))[:500]}"
        for r in all_results
    ])

    status_placeholder.info("🤖 AI extracting structured job listings...")
    claude_key = st.session_state.get("claude_key", "")
    raw = call_claude(
        "Extract job listings from web search results. Return ONLY valid JSON array, nothing else.",
        f'Extract all PM job listings. Return array: [{{"title":"","company":"","url":"","description":"","location":"","postedDate":""}}]\n\nContent:\n{content[:8000]}',
        claude_key
    )
    jobs = parse_json_safe(raw)
    return [j for j in (jobs or []) if j.get("title") and j.get("company")]


def stage_score(jobs: list, resume: str, status_placeholder) -> list:
    status_placeholder.info(f"🎯 Scoring {len(jobs)} jobs against your resume...")
    claude_key = st.session_state.get("claude_key", "")
    from concurrent.futures import ThreadPoolExecutor

    def score_job(j):
        raw = call_claude(
            "Score job match. Return ONLY valid JSON, nothing else.",
            f'Resume: {resume[:1200]}\n\nJob: {j["title"]} at {j["company"]}\nDesc: {str(j.get("description",""))[:700]}\n\nReturn: {{"score":8.5,"domain":9,"seniority":8,"technical":8,"ai_relevance":9,"match_reason":"2 sentences on why strong match","gap":"1 sentence or none","competition":"low/medium/high","sponsorship":"confirmed/likely/unknown/no","angle":"what to emphasize in application"}}',
            claude_key
        )
        parsed = parse_json_safe(raw)
        return {**j, **(parsed or {"score": 0})}

    scored = []
    with ThreadPoolExecutor(max_workers=4) as ex:
        results = list(ex.map(score_job, jobs[:20]))
        scored = results

    return sorted(
        [j for j in scored if j.get("score", 0) >= 7.5],
        key=lambda x: x.get("score", 0),
        reverse=True
    )[:8]


def stage_contacts(jobs: list, apify_key: str, status_placeholder) -> list:
    status_placeholder.info("👥 Finding VPs, Directors, Recruiters at each company...")
    claude_key = st.session_state.get("claude_key", "")
    result_jobs = []

    for i, job in enumerate(jobs):
        status_placeholder.info(f"👥 Finding decision makers at {job['company']} ({i+1}/{len(jobs)})...")
        try:
            results = run_apify(
                "apify/rag-web-browser",
                {
                    "query": f'"{job["company"]}" "VP Product" OR "Director Product" OR "Head of Product" OR "Technical Recruiter" site:linkedin.com',
                    "maxResults": 5
                },
                apify_key,
                timeout=120
            )
            content = "\n".join([f"{r.get('title','')} | {r.get('url','')}" for r in (results or [])])
            raw = call_claude(
                "Extract LinkedIn contacts. Return ONLY valid JSON array.",
                f'Find people at {job["company"]}. Return: [{{"name":"Full Name","title":"","linkedInUrl":"","priority":"high/medium/low"}}]\nPriority: VP/Director/Head Product=high, Recruiter=medium, PM=low\n\nResults:\n{content}',
                claude_key
            )
            contacts = parse_json_safe(raw)
            result_jobs.append({**job, "contacts": (contacts or [])[:5]})
        except Exception:
            result_jobs.append({**job, "contacts": []})

    return result_jobs


def stage_enrich_emails(jobs: list, hunter_key: str, status_placeholder) -> list:
    if not hunter_key:
        status_placeholder.warning("⚠️ No Hunter.io key — skipping email enrichment. Add it in the sidebar.")
        time.sleep(1)
        return jobs

    status_placeholder.info("📧 Looking up direct emails via Hunter.io...")
    found = 0
    total = 0
    enriched_jobs = []

    for job in jobs:
        enriched_contacts = []
        for contact in job.get("contacts", []):
            total += 1
            first, last = parse_name(contact.get("name", ""))
            result = hunter_find_email(first, last, job["company"], hunter_key)
            if result:
                found += 1
            status_placeholder.info(f"📧 Hunter.io: {found} emails found of {total} searched...")
            enriched_contacts.append({
                **contact,
                "email": result["email"] if result else None,
                "email_score": result["score"] if result else None,
                "email_verified": result["verified"] if result else False
            })
        enriched_jobs.append({**job, "contacts": enriched_contacts})

    status_placeholder.success(f"✅ Hunter.io complete: {found} direct emails found out of {total} contacts")
    return enriched_jobs


def stage_drafts(jobs: list, resume: str, status_placeholder) -> list:
    status_placeholder.info("✍️ Writing personalized outreach emails...")
    claude_key = st.session_state.get("claude_key", "")
    result_jobs = []

    for job in jobs:
        enriched_contacts = []
        for contact in job.get("contacts", []):
            try:
                draft = call_claude(
                    "Write a short cold outreach email. No hyphens. No bullet points. Max 4 sentences. Human and specific. No 'I am excited to'. Return email body only.",
                    f'Sender background: {resume[:350]}\nTo: {contact.get("name")}, {contact.get("title")} at {job["company"]}\nApplying for: {job["title"]}\n\nWrite email body:',
                    claude_key
                )
                enriched_contacts.append({**contact, "email_draft": draft.strip()})
            except Exception:
                enriched_contacts.append({**contact, "email_draft": ""})
        result_jobs.append({**job, "contacts": enriched_contacts})

    return result_jobs


def stage_resumes(jobs: list, resume: str, status_placeholder) -> list:
    status_placeholder.info("📄 Generating Resume Tailor prompts for each role...")
    result_jobs = []
    for job in jobs:
        prompt = f"""Tailor my resume for this role.

Company: {job.get("company", "")}
Role: {job.get("title", "")}
Apply Link: {job.get("url", "")}

Job Description:
{job.get("description", "Not available — paste JD here")}

Key angle to emphasize: {job.get("angle", "")}
Known gap to address: {job.get("gap", "")}

Instructions:
- Match keywords until score is above 90%
- Rewrite bullets using their exact language
- Flag any ATS keyword gaps
- Tell me what to add if a new gap is found
- Update memory with any new instructions"""

        result_jobs.append({{**job, "resume_tailor_prompt": prompt}})

    status_placeholder.success("✅ Resume Tailor prompts ready — paste each into your Resume Tailor project")
    return result_jobs


def generate_excel(jobs: list) -> bytes:
    """Generate formatted Excel file."""
    wb = openpyxl.Workbook()
    today = datetime.now().strftime("%Y-%m-%d")

    # ── Sheet 1: Jobs ──
    ws1 = wb.active
    ws1.title = "Jobs"
    headers1 = ["Date", "Rank", "Company", "Role", "Match Score", "Domain", "Seniority", "Technical", "AI Relevance",
                 "Apply Link", "Competition", "Sponsorship", "Why Match", "Address in Cover Note", "Tailoring Angle", "Status", "Notes"]
    ws1.append(headers1)
    for cell in ws1[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1a73e8")
        cell.alignment = Alignment(wrap_text=True)

    for i, job in enumerate(jobs):
        ws1.append([
            today, i+1, clean(job.get("company","")), clean(job.get("title","")),
            job.get("score",""), job.get("domain",""), job.get("seniority",""),
            job.get("technical",""), job.get("ai_relevance",""),
            clean(job.get("url","")), clean(job.get("competition","")), clean(job.get("sponsorship","")),
            clean(job.get("match_reason","")), clean(job.get("gap","")), clean(job.get("angle","")),
            "Not Applied", ""
        ])

    # ── Sheet 2: Contacts & Emails ──
    ws2 = wb.create_sheet("Contacts & Emails")
    headers2 = ["Date", "Company", "Role", "Contact Name", "Contact Title",
                 "Direct Email", "Email Confidence", "Email Verified", "LinkedIn URL",
                 "Priority", "Outreach Email Draft", "Sent?"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1a73e8")

    for job in jobs:
        for contact in job.get("contacts", []):
            ws2.append([
                today, clean(job.get("company","")), clean(job.get("title","")),
                clean(contact.get("name","")), clean(contact.get("title","")),
                clean(contact.get("email","Not found")),
                f'{contact.get("email_score","")}%' if contact.get("email_score") else "",
                "Yes" if contact.get("email_verified") else "",
                clean(contact.get("linkedInUrl","")),
                clean(contact.get("priority","")),
                clean(contact.get("email_draft","")),
                "No"
            ])

    # ── Sheet 3: Resume Tailor Prompts ──
    ws3 = wb.create_sheet("Resume Tailor Prompts")
    headers3 = ["Company", "Role", "Paste This Into Your Resume Tailor Project"]
    ws3.append(headers3)
    for cell in ws3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1a73e8")

    for job in jobs:
        ws3.append([clean(job.get("company","")), clean(job.get("title","")), clean(job.get("resume_tailor_prompt",""))])

    # Column widths
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ── Main UI ────────────────────────────────────────────────────────

def main():
    st.title("🚀 Daily Job Pipeline")
    st.caption(f"{datetime.now().strftime('%A, %B %d %Y')} · Goal: 10 interviews in 30 days")

    # ── Sidebar: Configuration ──
    with st.sidebar:
        st.header("⚙️ Configuration")
        st.caption("Keys are stored in your session only — never saved externally.")

        def get_secret(key, env_key):
            # Priority: session state → Streamlit secrets → env variable
            if st.session_state.get(key):
                return st.session_state[key]
            try:
                return st.secrets[env_key]
            except Exception:
                return os.environ.get(env_key, "")

        apify_key = st.text_input(
            "Apify API Key",
            type="password",
            value=get_secret("apify_key", "APIFY_KEY"),
            help="apify.com → Settings → API & Integrations"
        )
        claude_key = st.text_input(
            "Claude API Key",
            type="password",
            value=get_secret("claude_key", "ANTHROPIC_API_KEY"),
            help="console.anthropic.com → API Keys"
        )
        hunter_key = st.text_input(
            "Hunter.io API Key",
            type="password",
            value=get_secret("hunter_key", "HUNTER_KEY"),
            help="hunter.io → Dashboard → API Key · Optional but finds direct emails"
        )

        master_resume = st.text_area(
            "Master Resume",
            value=st.session_state.get("master_resume", ""),
            height=200,
            help="Paste your full resume text"
        )

        if st.button("💾 Save Configuration", use_container_width=True):
            st.session_state["apify_key"] = apify_key
            st.session_state["claude_key"] = claude_key
            st.session_state["hunter_key"] = hunter_key
            st.session_state["master_resume"] = master_resume
            st.success("Saved!")

        st.divider()
        st.caption("**Tips:**")
        st.caption("• Apify free tier works for testing")
        st.caption("• Claude API key from console.anthropic.com")
        st.caption("• Hunter.io free: 25 emails/month")
        st.caption("• Pipeline takes 10-15 mins to run")

        # Always visible download button in sidebar
        if st.session_state.get("excel_bytes"):
            st.divider()
            st.success("✅ Pipeline complete!")
            st.download_button(
                "📥 Download Excel Now",
                data=st.session_state["excel_bytes"],
                file_name=f"JobPipeline_{datetime.now().strftime('%Y-%m-%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="sidebar_download"
            )

    # ── Main: Pipeline ──
    if not st.session_state.get("apify_key") or not st.session_state.get("claude_key") or not st.session_state.get("master_resume"):
        st.warning("👈 Add your Apify key, Claude API key, and resume in the sidebar first. Then click Save Configuration.")
        return

    col1, col2 = st.columns([3, 1])
    with col1:
        st.subheader("Pipeline Stages")
    with col2:
        run_clicked = st.button("▶️ Run Pipeline", type="primary", use_container_width=True)

    # Stage display
    stages = [
        ("discover", "🔍", "Discover jobs"),
        ("score", "🎯", "Score matches"),
        ("contacts", "👥", "Find decision makers"),
        ("emails", "📧", "Enrich emails via Hunter.io"),
        ("drafts", "✍️", "Write outreach drafts"),
        ("resume", "📄", "Generate Resume Tailor prompts"),
        ("export", "📊", "Generate spreadsheet")
    ]

    done = st.session_state.get("done_stages", [])
    active = st.session_state.get("active_stage", None)

    cols = st.columns(4)
    for i, (sid, icon, label) in enumerate(stages):
        with cols[i % 4]:
            if sid in done:
                st.success(f"{icon} {label}")
            elif sid == active:
                st.info(f"{icon} {label} ⟳")
            else:
                st.markdown(f"<div style='color:#999;font-size:13px'>{icon} {label}</div>", unsafe_allow_html=True)

    status_area = st.empty()

    # ── Run Pipeline ──
    if run_clicked:
        st.session_state["done_stages"] = []
        st.session_state["active_stage"] = None
        st.session_state["pipeline_results"] = None

        try:
            ak = st.session_state["apify_key"]
            hk = st.session_state.get("hunter_key", "")
            resume = st.session_state["master_resume"]

            # Stage 1
            st.session_state["active_stage"] = "discover"
            jobs1 = stage_discover(ak, status_area)
            st.session_state["done_stages"] = ["discover"]
            status_area.success(f"✅ Found {len(jobs1)} job listings")

            # Stage 2
            st.session_state["active_stage"] = "score"
            jobs2 = stage_score(jobs1, resume, status_area)
            st.session_state["done_stages"] = ["discover", "score"]
            status_area.success(f"✅ {len(jobs2)} strong matches (score ≥ 7.5)")

            # Stage 3
            st.session_state["active_stage"] = "contacts"
            jobs3 = stage_contacts(jobs2, ak, status_area)
            st.session_state["done_stages"] = ["discover", "score", "contacts"]

            # Stage 4
            st.session_state["active_stage"] = "emails"
            jobs4 = stage_enrich_emails(jobs3, hk, status_area)
            st.session_state["done_stages"] = ["discover", "score", "contacts", "emails"]

            # Stage 5
            st.session_state["active_stage"] = "drafts"
            jobs5 = stage_drafts(jobs4, resume, status_area)
            st.session_state["done_stages"] = ["discover", "score", "contacts", "emails", "drafts"]

            # Stage 6
            st.session_state["active_stage"] = "resume"
            jobs6 = stage_resumes(jobs5, resume, status_area)
            st.session_state["done_stages"] = ["discover", "score", "contacts", "emails", "drafts", "resume"]

            # Stage 7
            st.session_state["active_stage"] = "export"
            excel_bytes = generate_excel(jobs6)
            st.session_state["done_stages"] = ["discover", "score", "contacts", "emails", "drafts", "resume", "export"]
            st.session_state["active_stage"] = None
            st.session_state["pipeline_results"] = jobs6
            st.session_state["excel_bytes"] = excel_bytes

            email_total = sum(1 for j in jobs6 for c in j.get("contacts", []) if c.get("email"))
            status_area.success(f"✅ Pipeline complete! {len(jobs6)} jobs · {email_total} direct emails · {len(jobs6)} Resume Tailor prompts ready")
            st.balloons()

        except Exception as e:
            status_area.error(f"❌ Pipeline error: {str(e)}")
            st.session_state["active_stage"] = None

    # ── Results ──
    results = st.session_state.get("pipeline_results")
    if results:
        st.divider()

        email_total = sum(1 for j in results for c in j.get("contacts", []) if c.get("email"))
        col1, col2, col3 = st.columns([2, 1, 1])
        with col1:
            st.subheader(f"✅ {len(results)} Matched Jobs · {email_total} Direct Emails Found")
        with col2:
            excel_bytes = st.session_state.get("excel_bytes")
            if excel_bytes:
                st.download_button(
                    "📥 Download Excel",
                    data=excel_bytes,
                    file_name=f"JobPipeline_{datetime.now().strftime('%Y-%m-%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

        # Job tabs
        tabs = st.tabs([f"{j.get('company','')} ({j.get('score',0):.1f})" for j in results])

        for tab, job in zip(tabs, results):
            with tab:
                col1, col2, col3 = st.columns(3)
                col1.metric("Match Score", f"{job.get('score',0):.1f}/10")
                col2.metric("Competition", job.get("competition","—").capitalize())
                col3.metric("Sponsorship", job.get("sponsorship","—").capitalize())

                if job.get("url"):
                    st.markdown(f"[🔗 Apply Now]({job['url']})")

                if job.get("match_reason"):
                    st.success(f"**Why you match:** {job['match_reason']}")

                if job.get("gap") and job["gap"].lower() != "none":
                    st.warning(f"**Address in cover note:** {job['gap']}")

                if job.get("resume_tailor_prompt"):
                    with st.expander("📋 Resume Tailor Prompt — Copy & paste into your Resume Tailor project"):
                        st.text_area("", value=job["resume_tailor_prompt"], height=200,
                            key=f"rtp_{job.get('company','')}",
                            label_visibility="collapsed",
                            help="Copy this and paste into your Resume Tailor Claude project")
                        st.caption("👆 Copy this → open your Resume Tailor project → paste → it will tailor to 90%+ match score automatically")

                st.subheader(f"👥 Decision Makers ({len(job.get('contacts',[]))})")
                for contact in job.get("contacts", []):
                    with st.expander(f"{'🌟' if contact.get('priority')=='high' else '🔵' if contact.get('priority')=='medium' else '⚪'} {contact.get('name','')} — {contact.get('title','')}"):
                        if contact.get("email"):
                            st.markdown(f"**📧 Direct Email:** `{contact['email']}` ({contact.get('email_score','')}% confidence{'  ✅ Verified' if contact.get('email_verified') else ''})")
                        else:
                            st.caption("📵 Email not found — use LinkedIn")

                        if contact.get("linkedInUrl"):
                            st.markdown(f"[LinkedIn Profile]({contact['linkedInUrl']})")

                        if contact.get("email_draft"):
                            st.markdown("**✍️ Outreach Email Draft:**")
                            st.text_area("", value=contact["email_draft"], height=120, key=f"draft_{job.get('company','')}_{contact.get('name','')}", label_visibility="collapsed")


if __name__ == "__main__":
    main()
